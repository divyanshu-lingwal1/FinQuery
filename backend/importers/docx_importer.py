from docx import Document


def read_docx_file(file_path):
    document = Document(file_path)

    text = ""

    # Read normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Read tables
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            text += " ".join(row_data) + "\n"

    return text

import re


def parse_docx_transactions(text):
    transactions = []

    pattern = re.compile(
        r"Date:\s*(.*?)\s*"
        r"Description:\s*(.*?)\s*"
        r"Category:\s*(.*?)\s*"
        r"Amount:\s*([\d,]+(?:\.\d+)?)",
        re.IGNORECASE
    )

    matches = pattern.findall(text)

    for match in matches:

        transaction = {
            "date": match[0].strip(),
            "description": match[1].strip(),
            "category": match[2].strip(),
            "amount": match[3].replace(",", "").strip()
        }

        transactions.append(transaction)

    return transactions